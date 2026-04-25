from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from app.models import AnalyticsEventRequest, Difficulty, ValidationRequest, ValidationResponse


LOG_DIRECTORY = Path(__file__).resolve().parents[2] / "data"
LOG_FILE_PATH = LOG_DIRECTORY / "validation_attempts.jsonl"
EVENT_LOG_FILE_PATH = LOG_DIRECTORY / "player_events.jsonl"
EXPORT_FILE_PATH = LOG_DIRECTORY / "skill_issue_player_logs.xlsx"


def log_validation_attempt(payload: ValidationRequest, result: ValidationResponse) -> str:
    LOG_DIRECTORY.mkdir(parents=True, exist_ok=True)
    attempt_id: str = str(uuid4())
    record = {
        "attempt_id": attempt_id,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "user_id": payload.user_id,
        "interaction_type": payload.interaction_type,
        "level_theme": payload.level_theme,
        "difficulty": payload.difficulty,
        "suggested_difficulty": result.suggested_difficulty,
        "time_taken_seconds": payload.time_taken_seconds,
        "success": result.success,
        "damage": result.damage,
        "errors": result.errors,
        "hints": result.hints,
        "resolved_effects": result.resolved_effects,
        "code_submitted": payload.code,
        "normalized_code": result.normalized_code,
    }
    with LOG_FILE_PATH.open("a", encoding="utf-8") as log_file:
        log_file.write(json.dumps(record, ensure_ascii=True) + "\n")
    return attempt_id


def log_player_event(payload: AnalyticsEventRequest) -> str:
    LOG_DIRECTORY.mkdir(parents=True, exist_ok=True)
    event_id: str = str(uuid4())
    timestamp_unix = payload.timestamp_unix
    if timestamp_unix <= 0:
        timestamp_unix = int(datetime.now(timezone.utc).timestamp())
    record = {
        "event_id": event_id,
        "timestamp": datetime.fromtimestamp(timestamp_unix, timezone.utc).isoformat(),
        "timestamp_unix": timestamp_unix,
        "user_id": payload.user_id,
        "session_id": payload.session_id,
        "save_slot": payload.save_slot,
        "event_type": payload.event_type,
        "stage_type": payload.stage_type,
        "level_id": payload.level_id,
        "progress_percent": payload.progress_percent,
        "metadata": payload.metadata,
    }
    with EVENT_LOG_FILE_PATH.open("a", encoding="utf-8") as log_file:
        log_file.write(json.dumps(record, ensure_ascii=True) + "\n")
    return event_id


def read_recent_attempts(limit: int = 20) -> list[dict]:
    if not LOG_FILE_PATH.exists():
        return []
    lines: list[str] = LOG_FILE_PATH.read_text(encoding="utf-8").splitlines()
    recent_lines: list[str] = lines[-max(limit, 1):]
    attempts: list[dict] = []
    for line in recent_lines:
        try:
            attempts.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return attempts


def export_player_logs_workbook() -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    LOG_DIRECTORY.mkdir(parents=True, exist_ok=True)
    attempts = _read_jsonl(LOG_FILE_PATH)
    events = _read_jsonl(EVENT_LOG_FILE_PATH)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    attempts_sheet = workbook.create_sheet("Validation Attempts")
    events_sheet = workbook.create_sheet("Player Events")
    event_counts_sheet = workbook.create_sheet("Event Counts")
    timeline_sheet = workbook.create_sheet("Timeline")

    _write_summary(summary_sheet, attempts, events)
    _write_table(attempts_sheet, attempts, [
        "attempt_id",
        "timestamp",
        "user_id",
        "interaction_type",
        "level_theme",
        "difficulty",
        "suggested_difficulty",
        "time_taken_seconds",
        "success",
        "damage",
        "errors",
        "hints",
        "resolved_effects",
        "code_submitted",
        "normalized_code",
    ])
    _write_table(events_sheet, events, [
        "event_id",
        "timestamp",
        "user_id",
        "session_id",
        "save_slot",
        "event_type",
        "stage_type",
        "level_id",
        "progress_percent",
        "metadata",
    ])
    _write_event_counts(event_counts_sheet, events)
    _write_timeline(timeline_sheet, events, attempts)

    _add_bar_chart(event_counts_sheet, BarChart, Reference)
    _add_timeline_chart(timeline_sheet, LineChart, Reference)
    for sheet in workbook.worksheets:
        _style_sheet(sheet, Font, PatternFill, get_column_letter)

    workbook.save(EXPORT_FILE_PATH)
    return EXPORT_FILE_PATH


def summarize_user_theme_attempts(user_id: str, level_theme: str, limit: int = 12) -> dict[str, int | Difficulty]:
    if not LOG_FILE_PATH.exists():
        return {
            "successes": 0,
            "failures": 0,
            "current_streak": 0,
            "recent_count": 0,
            "suggested_difficulty": "easy",
        }

    attempts: list[dict] = []
    lines: list[str] = LOG_FILE_PATH.read_text(encoding="utf-8").splitlines()
    for raw_line in reversed(lines):
        try:
            record = json.loads(raw_line)
        except json.JSONDecodeError:
            continue
        if record.get("user_id") != user_id or record.get("level_theme") != level_theme:
            continue
        attempts.append(record)
        if len(attempts) >= max(limit, 1):
            break

    if not attempts:
        return {
            "successes": 0,
            "failures": 0,
            "current_streak": 0,
            "recent_count": 0,
            "suggested_difficulty": "easy",
        }

    successes: int = sum(1 for item in attempts if item.get("success"))
    failures: int = len(attempts) - successes
    current_streak: int = 0
    streak_sign: int = 0
    for item in attempts:
        item_success: bool = bool(item.get("success"))
        item_sign: int = 1 if item_success else -1
        if streak_sign == 0 or item_sign == streak_sign:
            current_streak += item_sign
            streak_sign = item_sign
            continue
        break

    suggested_difficulty: Difficulty = "easy"
    if successes >= failures + 2:
        suggested_difficulty = "normal"
    if successes >= failures + 5:
        suggested_difficulty = "hard"
    if current_streak <= -3:
        suggested_difficulty = "easy"

    return {
        "successes": successes,
        "failures": failures,
        "current_streak": current_streak,
        "recent_count": len(attempts),
        "suggested_difficulty": suggested_difficulty,
    }


def _read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    records: list[dict] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            value = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(value, dict):
            records.append(value)
    return records


def _write_summary(sheet, attempts: list[dict], events: list[dict]) -> None:
    success_count = sum(1 for attempt in attempts if attempt.get("success"))
    failure_count = len(attempts) - success_count
    event_types = {str(event.get("event_type", "")) for event in events if event.get("event_type")}
    terminal_opens = sum(1 for event in events if event.get("event_type") == "terminal_opened")
    hint_opens = sum(1 for event in events if event.get("event_type") == "hint_requested")
    explanation_opens = sum(1 for event in events if event.get("event_type") == "explanation_opened")
    completions = sum(1 for event in events if str(event.get("event_type", "")).endswith("_completed"))
    success_rate = round(success_count / len(attempts) * 100, 1) if attempts else 0

    rows = [
        ("Metric", "Value"),
        ("Total player events", len(events)),
        ("Unique event types", len(event_types)),
        ("Terminal opens", terminal_opens),
        ("Hint opens", hint_opens),
        ("Explanation opens", explanation_opens),
        ("Stage completions", completions),
        ("Validation attempts", len(attempts)),
        ("Successful validations", success_count),
        ("Failed validations", failure_count),
        ("Validation success rate, %", success_rate),
        ("Generated at UTC", datetime.now(timezone.utc).isoformat()),
    ]
    for row in rows:
        sheet.append(row)


def _write_table(sheet, records: list[dict], headers: list[str]) -> None:
    sheet.append(headers)
    for record in records:
        sheet.append([_cell_value(record.get(header, "")) for header in headers])


def _write_event_counts(sheet, events: list[dict]) -> None:
    counts: dict[str, int] = {}
    for event in events:
        event_type = str(event.get("event_type", "unknown"))
        counts[event_type] = counts.get(event_type, 0) + 1
    sheet.append(["event_type", "count"])
    for event_type, count in sorted(counts.items(), key=lambda item: item[1], reverse=True):
        sheet.append([event_type, count])
    if not counts:
        sheet.append(["no_events_yet", 0])


def _write_timeline(sheet, events: list[dict], attempts: list[dict]) -> None:
    buckets: dict[str, dict[str, int]] = {}
    for event in events:
        bucket = _hour_bucket(str(event.get("timestamp", "")))
        buckets.setdefault(bucket, {"events": 0, "attempts": 0, "successes": 0})
        buckets[bucket]["events"] += 1
    for attempt in attempts:
        bucket = _hour_bucket(str(attempt.get("timestamp", "")))
        buckets.setdefault(bucket, {"events": 0, "attempts": 0, "successes": 0})
        buckets[bucket]["attempts"] += 1
        if attempt.get("success"):
            buckets[bucket]["successes"] += 1
    sheet.append(["hour_utc", "events", "validation_attempts", "successful_validations"])
    for bucket in sorted(buckets.keys()):
        sheet.append([bucket, buckets[bucket]["events"], buckets[bucket]["attempts"], buckets[bucket]["successes"]])
    if not buckets:
        sheet.append([datetime.now(timezone.utc).strftime("%Y-%m-%d %H:00"), 0, 0, 0])


def _add_bar_chart(sheet, bar_chart_cls, reference_cls) -> None:
    if sheet.max_row < 2:
        return
    chart = bar_chart_cls()
    chart.title = "Player Event Counts"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Event Type"
    data = reference_cls(sheet, min_col=2, min_row=1, max_row=sheet.max_row)
    categories = reference_cls(sheet, min_col=1, min_row=2, max_row=sheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 16
    sheet.add_chart(chart, "D2")


def _add_timeline_chart(sheet, line_chart_cls, reference_cls) -> None:
    if sheet.max_row < 2:
        return
    chart = line_chart_cls()
    chart.title = "Activity Timeline"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Hour UTC"
    data = reference_cls(sheet, min_col=2, max_col=4, min_row=1, max_row=sheet.max_row)
    categories = reference_cls(sheet, min_col=1, min_row=2, max_row=sheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 18
    sheet.add_chart(chart, "F2")


def _style_sheet(sheet, font_cls, fill_cls, column_letter_fn) -> None:
    header_fill = fill_cls("solid", fgColor="2D2140")
    header_font = font_cls(bold=True, color="FFD45E")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for column_index in range(1, sheet.max_column + 1):
        column_letter = column_letter_fn(column_index)
        max_length = 12
        for cell in sheet[column_letter]:
            max_length = max(max_length, min(len(str(cell.value or "")), 72))
        sheet.column_dimensions[column_letter].width = max_length + 2


def _cell_value(value) -> str | int | float | bool:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return "" if value is None else value
    return json.dumps(value, ensure_ascii=False)


def _hour_bucket(timestamp: str) -> str:
    try:
        parsed = datetime.fromisoformat(timestamp.replace("Z", "+00:00"))
    except ValueError:
        parsed = datetime.now(timezone.utc)
    return parsed.astimezone(timezone.utc).strftime("%Y-%m-%d %H:00")
