import json
import sys
from pathlib import Path

from docx import Document


def para_style_name(p):
    try:
        return p.style.name or ""
    except Exception:
        return ""


def read_docx(path):
    doc = Document(path)
    paragraphs = []
    headings = []
    for i, p in enumerate(doc.paragraphs):
        text = " ".join((p.text or "").split())
        if not text:
            continue
        style = para_style_name(p)
        item = {"index": i, "style": style, "text": text}
        paragraphs.append(item)
        if "Heading" in style or "Заголов" in style:
            headings.append(item)
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows[:8]:
            rows.append([" ".join((cell.text or "").split()) for cell in row.cells])
        tables.append({"index": ti, "rows": len(table.rows), "cols": len(table.columns), "sample": rows})
    return {"paragraph_count": len(paragraphs), "headings": headings, "paragraphs": paragraphs, "tables": tables}


def read_pdf(path):
    data = {"pages": []}
    try:
        import fitz
        pdf = fitz.open(path)
        data["page_count"] = len(pdf)
        for idx, page in enumerate(pdf):
            text = " ".join((page.get_text("text") or "").split())
            data["pages"].append({"page": idx + 1, "text": text[:2500]})
    except Exception as exc:
        data["fitz_error"] = str(exc)
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            data["page_count"] = len(reader.pages)
            for idx, page in enumerate(reader.pages):
                text = " ".join((page.extract_text() or "").split())
                data["pages"].append({"page": idx + 1, "text": text[:2500]})
        except Exception as exc2:
            data["error"] = f"pypdf failed: {exc2}"
    return data


def main():
    docx_path = Path(sys.argv[1])
    pdf_path = Path(sys.argv[2])
    out_path = Path(sys.argv[3])
    result = {
        "docx": read_docx(docx_path),
        "pdf": read_pdf(pdf_path),
    }
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
