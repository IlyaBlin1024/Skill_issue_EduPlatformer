from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from app.models import AnalyticsEventRequest, Difficulty, ValidationRequest, ValidationResponse


LOG_DIRECTORY = Path(__file__).resolve().parents[2] / "data"
LOG_FILE_PATH = LOG_DIRECTORY / "validation_attempts.jsonl"
EVENT_LOG_FILE_PATH = LOG_DIRECTORY / "player_events.jsonl"
GENERATED_TASKS_FILE_PATH = LOG_DIRECTORY / "generated_tasks.jsonl"
EXPORT_FILE_PATH = LOG_DIRECTORY / "skill_issue_player_logs.xlsx"


def log_validation_attempt(payload: ValidationRequest, result: ValidationResponse) -> str:
    LOG_DIRECTORY.mkdir(parents=True, exist_ok=True)
    attempt_id: str = str(uuid4())
    record = {
        "attempt_id": attempt_id,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "user_id": payload.user_id,
        "interaction_type": payload.interaction_type,
        "stage_type": payload.stage_type,
        "level_id": payload.level_id,
        "level_theme": payload.level_theme,
        "difficulty": payload.difficulty,
        "suggested_difficulty": result.suggested_difficulty,
        "time_taken_seconds": payload.time_taken_seconds,
        "success": result.success,
        "damage": result.damage,
        "errors": result.errors,
        "hints": result.hints,
        "resolved_effects": result.resolved_effects,
        "generation_source": payload.generation_source,
        "generation_detail": payload.generation_detail,
        "pattern_id": payload.pattern_id,
        "terminal_session_id": payload.terminal_session_id,
        "attempt_number": payload.attempt_number,
        "task_generation_seconds": payload.task_generation_seconds,
        "task_prompt": payload.task_prompt,
        "keywords": payload.keywords,
        "validation_targets": payload.validation_targets,
        "hint_used": payload.hint_used,
        "explanation_used": payload.explanation_used,
        "hint_request_count": payload.hint_request_count,
        "explanation_open_count": payload.explanation_open_count,
        "explanation_read_seconds": payload.explanation_read_seconds,
        "code_submitted": payload.code,
        "normalized_code": result.normalized_code,
    }
    with LOG_FILE_PATH.open("a", encoding="utf-8") as log_file:
        log_file.write(json.dumps(record, ensure_ascii=True) + "\n")
    return attempt_id


def log_player_event(payload: AnalyticsEventRequest) -> str:
    LOG_DIRECTORY.mkdir(parents=True, exist_ok=True)
    event_id: str = str(uuid4())
    timestamp_unix = payload.timestamp_unix
    if timestamp_unix <= 0:
        timestamp_unix = int(datetime.now(timezone.utc).timestamp())
    record = {
        "event_id": event_id,
        "timestamp": datetime.fromtimestamp(timestamp_unix, timezone.utc).isoformat(),
        "timestamp_unix": timestamp_unix,
        "user_id": payload.user_id,
        "session_id": payload.session_id,
        "save_slot": payload.save_slot,
        "event_type": payload.event_type,
        "stage_type": payload.stage_type,
        "level_id": payload.level_id,
        "progress_percent": payload.progress_percent,
        "metadata": payload.metadata,
    }
    with EVENT_LOG_FILE_PATH.open("a", encoding="utf-8") as log_file:
        log_file.write(json.dumps(record, ensure_ascii=True) + "\n")
    return event_id


def read_recent_attempts(limit: int = 20) -> list[dict]:
    if not LOG_FILE_PATH.exists():
        return []
    lines: list[str] = LOG_FILE_PATH.read_text(encoding="utf-8").splitlines()
    recent_lines: list[str] = lines[-max(limit, 1):]
    attempts: list[dict] = []
    for line in recent_lines:
        try:
            attempts.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return attempts


def export_player_logs_workbook() -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    LOG_DIRECTORY.mkdir(parents=True, exist_ok=True)
    attempts = _read_jsonl(LOG_FILE_PATH)
    events = _read_jsonl(EVENT_LOG_FILE_PATH)
    generated_tasks = _read_jsonl(GENERATED_TASKS_FILE_PATH)
    datasets = _build_export_datasets(attempts, events, generated_tasks)

    workbook = Workbook()
    overview_sheet = workbook.active
    overview_sheet.title = "Overview"
    conclusions_sheet = workbook.create_sheet("Conclusions")
    learning_sheet = workbook.create_sheet("Learning Metrics")
    theme_sheet = workbook.create_sheet("Theme Metrics")
    interaction_sheet = workbook.create_sheet("Interaction Metrics")
    support_sheet = workbook.create_sheet("Support Impact")
    errors_sheet = workbook.create_sheet("Error Analysis")
    ai_sheet = workbook.create_sheet("AI Quality")
    terminal_sheet = workbook.create_sheet("Terminal Tasks")
    behavior_sheet = workbook.create_sheet("Game Behavior")
    difficulty_sheet = workbook.create_sheet("Difficulty Progression")
    funnel_sheet = workbook.create_sheet("Progression Funnel")
    boss_sheet = workbook.create_sheet("Boss Performance")
    retention_sheet = workbook.create_sheet("Retention")
    timeline_sheet = workbook.create_sheet("Timeline")
    event_counts_sheet = workbook.create_sheet("Event Counts")
    charts_sheet = workbook.create_sheet("Charts")
    attempts_sheet = workbook.create_sheet("Validation Attempts")
    events_sheet = workbook.create_sheet("Player Events")
    generated_tasks_sheet = workbook.create_sheet("Generated Tasks")

    _write_overview(overview_sheet, datasets["overview"], datasets["research_questions"])
    _write_conclusions(conclusions_sheet, datasets["conclusions"])
    _write_key_value_table(learning_sheet, "Learning Metrics", datasets["learning_metrics"])
    _write_table(theme_sheet, datasets["theme_metrics"], [
        "level_theme",
        "attempts",
        "successes",
        "failures",
        "success_rate_percent",
        "avg_time_seconds",
        "hint_used_attempts",
        "explanation_used_attempts",
        "hint_events",
        "explanation_events",
        "ai_tasks",
        "fallback_tasks",
    ])
    _write_table(interaction_sheet, datasets["interaction_metrics"], [
        "interaction_type",
        "attempts",
        "successes",
        "failures",
        "success_rate_percent",
        "avg_time_seconds",
        "hint_used_attempts",
        "explanation_used_attempts",
    ])
    _write_table(support_sheet, datasets["support_impact"], [
        "support_type",
        "attempts",
        "successes",
        "success_rate_percent",
        "avg_time_seconds",
    ])
    _write_table(errors_sheet, datasets["error_analysis"], [
        "error",
        "count",
        "most_common_theme",
        "most_common_interaction",
    ])
    _write_table(ai_sheet, datasets["ai_quality"], [
        "category",
        "count",
        "share_percent",
        "detail",
    ])
    _write_table(terminal_sheet, datasets["terminal_tasks"], [
        "terminal_session_id",
        "level_theme",
        "interaction_type",
        "attempts",
        "success",
        "first_try_success",
        "time_to_success_seconds",
        "hint_used",
        "explanation_used",
        "explanation_read_seconds",
        "generation_source",
    ])
    _write_table(behavior_sheet, datasets["game_behavior"], [
        "behavior",
        "count",
        "total_value",
        "average_value",
        "interpretation",
    ])
    _write_table(difficulty_sheet, datasets["difficulty_progression"], [
        "timestamp",
        "interaction_type",
        "level_theme",
        "from_difficulty",
        "to_difficulty",
        "reason",
    ])
    _write_table(funnel_sheet, datasets["progression_funnel"], [
        "stage",
        "started",
        "completed",
        "completion_rate_percent",
        "dropoff_count",
    ])
    _write_table(boss_sheet, datasets["boss_performance"], [
        "level_id",
        "attempts",
        "victories",
        "deaths",
        "terminal_rewrites",
        "player_damage_taken",
        "damage_dealt_to_boss",
        "parry_successes",
    ])
    _write_table(retention_sheet, datasets["retention"], [
        "date_utc",
        "sessions",
        "events",
        "returning_session_count",
        "session_duration_seconds",
    ])
    _write_table(timeline_sheet, datasets["timeline"], [
        "hour_utc",
        "events",
        "validation_attempts",
        "successful_validations",
        "failed_validations",
        "hint_requests",
        "explanation_opens",
    ])
    _write_table(event_counts_sheet, datasets["event_counts"], ["event_type", "count"])
    _write_table(attempts_sheet, attempts, [
        "attempt_id",
        "timestamp",
        "user_id",
        "interaction_type",
        "stage_type",
        "level_id",
        "level_theme",
        "difficulty",
        "suggested_difficulty",
        "time_taken_seconds",
        "success",
        "damage",
        "errors",
        "hints",
        "resolved_effects",
        "generation_source",
        "generation_detail",
        "pattern_id",
        "terminal_session_id",
        "attempt_number",
        "task_generation_seconds",
        "task_prompt",
        "keywords",
        "validation_targets",
        "hint_used",
        "explanation_used",
        "hint_request_count",
        "explanation_open_count",
        "explanation_read_seconds",
        "code_submitted",
        "normalized_code",
    ])
    _write_table(events_sheet, events, [
        "event_id",
        "timestamp",
        "user_id",
        "session_id",
        "save_slot",
        "event_type",
        "stage_type",
        "level_id",
        "progress_percent",
        "metadata",
    ])
    _write_table(generated_tasks_sheet, generated_tasks, [
        "timestamp",
        "storage_mode",
        "interaction_type",
        "level_theme",
        "difficulty",
        "generation_source",
        "generation_detail",
        "pattern_id",
        "title",
        "prompt",
        "keywords",
        "validation_targets",
    ])

    _add_research_charts(
        charts_sheet,
        event_counts_sheet,
        timeline_sheet,
        theme_sheet,
        interaction_sheet,
        errors_sheet,
        ai_sheet,
        funnel_sheet,
        behavior_sheet,
        boss_sheet,
        BarChart,
        LineChart,
        PieChart,
        Reference,
    )
    for sheet in workbook.worksheets:
        _style_sheet(sheet, Font, PatternFill, Alignment, Border, Side, get_column_letter)
        _add_table_if_possible(sheet, Table, TableStyleInfo)

    workbook.save(EXPORT_FILE_PATH)
    return EXPORT_FILE_PATH


def summarize_user_theme_attempts(user_id: str, level_theme: str, limit: int = 12) -> dict[str, int | Difficulty]:
    if not LOG_FILE_PATH.exists():
        return {
            "successes": 0,
            "failures": 0,
            "current_streak": 0,
            "recent_count": 0,
            "suggested_difficulty": "easy",
        }

    attempts: list[dict] = []
    lines: list[str] = LOG_FILE_PATH.read_text(encoding="utf-8").splitlines()
    for raw_line in reversed(lines):
        try:
            record = json.loads(raw_line)
        except json.JSONDecodeError:
            continue
        if record.get("user_id") != user_id or record.get("level_theme") != level_theme:
            continue
        attempts.append(record)
        if len(attempts) >= max(limit, 1):
            break

    if not attempts:
        return {
            "successes": 0,
            "failures": 0,
            "current_streak": 0,
            "recent_count": 0,
            "suggested_difficulty": "easy",
        }

    successes: int = sum(1 for item in attempts if item.get("success"))
    failures: int = len(attempts) - successes
    current_streak: int = 0
    streak_sign: int = 0
    for item in attempts:
        item_success: bool = bool(item.get("success"))
        item_sign: int = 1 if item_success else -1
        if streak_sign == 0 or item_sign == streak_sign:
            current_streak += item_sign
            streak_sign = item_sign
            continue
        break

    suggested_difficulty: Difficulty = "easy"
    if successes >= failures + 2:
        suggested_difficulty = "normal"
    if successes >= failures + 5:
        suggested_difficulty = "hard"
    if current_streak <= -3:
        suggested_difficulty = "easy"

    return {
        "successes": successes,
        "failures": failures,
        "current_streak": current_streak,
        "recent_count": len(attempts),
        "suggested_difficulty": suggested_difficulty,
    }


def _read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    records: list[dict] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            value = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(value, dict):
            records.append(value)
    return records


def _build_export_datasets(attempts: list[dict], events: list[dict], generated_tasks: list[dict]) -> dict[str, list[dict]]:
    success_count = sum(1 for attempt in attempts if attempt.get("success"))
    failure_count = len(attempts) - success_count
    terminal_opens = _count_events(events, "terminal_opened")
    hint_opens = _count_events(events, "hint_requested")
    explanation_opens = _count_events(events, "explanation_opened")
    completions = sum(1 for event in events if str(event.get("event_type", "")).endswith("_completed"))
    deaths = _count_events(events, "player_died") + _count_events(events, "death")
    sessions = len({str(event.get("session_id", "")) for event in events if event.get("session_id")})
    ai_tasks = sum(1 for task in generated_tasks if str(task.get("generation_source", "")).lower() == "ai")
    fallback_tasks = sum(1 for task in generated_tasks if str(task.get("generation_source", "")).lower() == "fallback")
    prewarm_failures = sum(1 for task in generated_tasks if str(task.get("storage_mode", "")) == "prewarm_failed")
    success_rate = _percent(success_count, len(attempts))
    avg_time = _average([_safe_float(attempt.get("time_taken_seconds", 0)) for attempt in attempts])
    avg_success_time = _average([_safe_float(attempt.get("time_taken_seconds", 0)) for attempt in attempts if attempt.get("success")])
    hint_attempts = sum(1 for attempt in attempts if bool(attempt.get("hint_used")))
    explanation_attempts = sum(1 for attempt in attempts if bool(attempt.get("explanation_used")))
    terminal_tasks = _terminal_tasks(attempts)
    solved_tasks = [task for task in terminal_tasks if bool(task.get("success"))]
    first_try_successes = sum(1 for task in solved_tasks if bool(task.get("first_try_success")))

    learning_metrics = [
        {"metric": "Player events", "value": len(events), "interpretation": "Raw behavioral events collected from the game."},
        {"metric": "Sessions", "value": sessions, "interpretation": "Distinct play sessions in the exported logs."},
        {"metric": "Terminal opens", "value": terminal_opens, "interpretation": "How often players reached a coding task."},
        {"metric": "Validation attempts", "value": len(attempts), "interpretation": "Submitted code snippets checked by backend."},
        {"metric": "Successful validations", "value": success_count, "interpretation": "Accepted solutions."},
        {"metric": "Failed validations", "value": failure_count, "interpretation": "Rejected solutions."},
        {"metric": "Validation success rate, %", "value": success_rate, "interpretation": "Primary learning-performance signal."},
        {"metric": "First-try success rate, %", "value": _percent(first_try_successes, len(terminal_tasks)), "interpretation": "Whether the player can solve a fresh task without repeated correction."},
        {"metric": "Average attempts to success", "value": _average([_safe_float(task.get("attempts", 0)) for task in solved_tasks]), "interpretation": "How many code submissions successful tasks require."},
        {"metric": "Average time per attempt, sec", "value": avg_time, "interpretation": "General task effort and friction."},
        {"metric": "Average successful time, sec", "value": avg_success_time, "interpretation": "How long solved tasks take."},
        {"metric": "Hint requests", "value": hint_opens, "interpretation": "Momentary support demand."},
        {"metric": "Explanation opens", "value": explanation_opens, "interpretation": "Concept support demand."},
        {"metric": "Attempts after hint, %", "value": _percent(hint_attempts, len(attempts)), "interpretation": "How dependent solutions are on hints."},
        {"metric": "Attempts after explanation, %", "value": _percent(explanation_attempts, len(attempts)), "interpretation": "How often players study before solving."},
        {"metric": "Average explanation read time, sec", "value": _average([_safe_float(attempt.get("explanation_read_seconds", 0)) for attempt in attempts]), "interpretation": "How long players spend reading topic support before submitting."},
        {"metric": "Stage completions", "value": completions, "interpretation": "Macro-progression through tutorial, levels, and bosses."},
        {"metric": "Deaths", "value": deaths, "interpretation": "Game difficulty and frustration proxy."},
        {"metric": "AI-generated tasks", "value": ai_tasks, "interpretation": "Live AI content delivered."},
        {"metric": "Fallback pattern tasks", "value": fallback_tasks, "interpretation": "Static backup content delivered."},
        {"metric": "AI prewarm failures", "value": prewarm_failures, "interpretation": "AI provider or parsing failures."},
        {"metric": "Generated at UTC", "value": datetime.now(timezone.utc).isoformat(), "interpretation": "Export timestamp."},
    ]

    datasets = {
        "overview": learning_metrics,
        "research_questions": _research_questions(),
        "learning_metrics": learning_metrics,
        "theme_metrics": _group_attempt_metrics(attempts, events, generated_tasks, "level_theme"),
        "interaction_metrics": _group_attempt_metrics(attempts, events, generated_tasks, "interaction_type"),
        "support_impact": _support_impact(attempts),
        "error_analysis": _error_analysis(attempts),
        "ai_quality": _ai_quality(generated_tasks),
        "terminal_tasks": _terminal_tasks(attempts),
        "game_behavior": _game_behavior(events),
        "difficulty_progression": _difficulty_progression(events),
        "progression_funnel": _progression_funnel(events),
        "boss_performance": _boss_performance(attempts, events),
        "retention": _retention(events),
        "timeline": _timeline(events, attempts),
        "event_counts": _event_counts(events),
        "conclusions": _build_conclusions(learning_metrics, attempts, events, generated_tasks),
    }
    return datasets


def _write_overview(sheet, metrics: list[dict], research_questions: list[dict]) -> None:
    sheet.append(["Skill Issue Learning Impact Export", ""])
    sheet.append(["Generated", datetime.now(timezone.utc).isoformat()])
    sheet.append([])
    sheet.append(["Metric", "Value", "Interpretation"])
    for record in metrics:
        sheet.append([record["metric"], record["value"], record["interpretation"]])
    sheet.append([])
    sheet.append(["Research Question", "How to read it"])
    for record in research_questions:
        sheet.append([record["question"], record["reading"]])


def _write_conclusions(sheet, conclusions: list[dict]) -> None:
    sheet.append(["Conclusion", "Signal", "Recommended action"])
    for record in conclusions:
        sheet.append([record["conclusion"], record["signal"], record["recommended_action"]])


def _write_key_value_table(sheet, title: str, records: list[dict]) -> None:
    sheet.append([title, "", ""])
    sheet.append(["Metric", "Value", "Interpretation"])
    for record in records:
        sheet.append([record["metric"], record["value"], record["interpretation"]])


def _write_table(sheet, records: list[dict], headers: list[str]) -> None:
    sheet.append(headers)
    for record in records:
        sheet.append([_cell_value(record.get(header, "")) for header in headers])
    if not records:
        sheet.append(["" for _header in headers])


def _add_research_charts(
    charts_sheet,
    event_counts_sheet,
    timeline_sheet,
    theme_sheet,
    interaction_sheet,
    errors_sheet,
    ai_sheet,
    funnel_sheet,
    behavior_sheet,
    boss_sheet,
    bar_chart_cls,
    line_chart_cls,
    pie_chart_cls,
    reference_cls,
) -> None:
    charts_sheet.append(["Charts"])
    _add_chart_title(charts_sheet, "A2", "Activity and learning overview")
    if event_counts_sheet.max_row >= 2:
        event_chart = bar_chart_cls()
        event_chart.title = "Player Event Counts"
        event_chart.y_axis.title = "Count"
        event_chart.x_axis.title = "Event Type"
        event_chart.add_data(reference_cls(event_counts_sheet, min_col=2, min_row=1, max_row=event_counts_sheet.max_row), titles_from_data=True)
        event_chart.set_categories(reference_cls(event_counts_sheet, min_col=1, min_row=2, max_row=event_counts_sheet.max_row))
        event_chart.height = 8
        event_chart.width = 14
        charts_sheet.add_chart(event_chart, "A4")
    if timeline_sheet.max_row >= 2:
        timeline_chart = line_chart_cls()
        timeline_chart.title = "Learning Timeline"
        timeline_chart.y_axis.title = "Count"
        timeline_chart.x_axis.title = "Hour UTC"
        timeline_chart.add_data(reference_cls(timeline_sheet, min_col=2, max_col=7, min_row=1, max_row=timeline_sheet.max_row), titles_from_data=True)
        timeline_chart.set_categories(reference_cls(timeline_sheet, min_col=1, min_row=2, max_row=timeline_sheet.max_row))
        timeline_chart.height = 8
        timeline_chart.width = 18
        charts_sheet.add_chart(timeline_chart, "P4")
    if theme_sheet.max_row >= 2:
        theme_chart = bar_chart_cls()
        theme_chart.title = "Success Rate by Theme"
        theme_chart.y_axis.title = "Success Rate, %"
        theme_chart.add_data(reference_cls(theme_sheet, min_col=5, min_row=1, max_row=theme_sheet.max_row), titles_from_data=True)
        theme_chart.set_categories(reference_cls(theme_sheet, min_col=1, min_row=2, max_row=theme_sheet.max_row))
        theme_chart.height = 8
        theme_chart.width = 14
        charts_sheet.add_chart(theme_chart, "A22")
    if interaction_sheet.max_row >= 2:
        interaction_chart = bar_chart_cls()
        interaction_chart.title = "Success Rate by Interaction"
        interaction_chart.y_axis.title = "Success Rate, %"
        interaction_chart.add_data(reference_cls(interaction_sheet, min_col=5, min_row=1, max_row=interaction_sheet.max_row), titles_from_data=True)
        interaction_chart.set_categories(reference_cls(interaction_sheet, min_col=1, min_row=2, max_row=interaction_sheet.max_row))
        interaction_chart.height = 8
        interaction_chart.width = 14
        charts_sheet.add_chart(interaction_chart, "P22")
    if errors_sheet.max_row >= 2:
        error_chart = bar_chart_cls()
        error_chart.title = "Most Common Validation Errors"
        error_chart.y_axis.title = "Count"
        error_chart.add_data(reference_cls(errors_sheet, min_col=2, min_row=1, max_row=min(errors_sheet.max_row, 12)), titles_from_data=True)
        error_chart.set_categories(reference_cls(errors_sheet, min_col=1, min_row=2, max_row=min(errors_sheet.max_row, 12)))
        error_chart.height = 8
        error_chart.width = 18
        charts_sheet.add_chart(error_chart, "A40")
    if ai_sheet.max_row >= 2:
        ai_chart = pie_chart_cls()
        ai_chart.title = "AI / Fallback Content Mix"
        ai_chart.add_data(reference_cls(ai_sheet, min_col=2, min_row=1, max_row=ai_sheet.max_row), titles_from_data=True)
        ai_chart.set_categories(reference_cls(ai_sheet, min_col=1, min_row=2, max_row=ai_sheet.max_row))
        ai_chart.height = 8
        ai_chart.width = 12
        charts_sheet.add_chart(ai_chart, "P40")
    if funnel_sheet.max_row >= 2:
        funnel_chart = bar_chart_cls()
        funnel_chart.title = "Progression Funnel"
        funnel_chart.y_axis.title = "Count"
        funnel_chart.add_data(reference_cls(funnel_sheet, min_col=2, max_col=3, min_row=1, max_row=funnel_sheet.max_row), titles_from_data=True)
        funnel_chart.set_categories(reference_cls(funnel_sheet, min_col=1, min_row=2, max_row=funnel_sheet.max_row))
        funnel_chart.height = 8
        funnel_chart.width = 18
        charts_sheet.add_chart(funnel_chart, "A58")
    if behavior_sheet.max_row >= 2:
        behavior_chart = bar_chart_cls()
        behavior_chart.title = "Combat and Engagement Behavior"
        behavior_chart.y_axis.title = "Count"
        behavior_chart.add_data(reference_cls(behavior_sheet, min_col=2, min_row=1, max_row=behavior_sheet.max_row), titles_from_data=True)
        behavior_chart.set_categories(reference_cls(behavior_sheet, min_col=1, min_row=2, max_row=behavior_sheet.max_row))
        behavior_chart.height = 8
        behavior_chart.width = 18
        charts_sheet.add_chart(behavior_chart, "P58")
    if boss_sheet.max_row >= 2:
        boss_chart = bar_chart_cls()
        boss_chart.title = "Boss Performance"
        boss_chart.y_axis.title = "Count"
        boss_chart.add_data(reference_cls(boss_sheet, min_col=2, max_col=5, min_row=1, max_row=boss_sheet.max_row), titles_from_data=True)
        boss_chart.set_categories(reference_cls(boss_sheet, min_col=1, min_row=2, max_row=boss_sheet.max_row))
        boss_chart.height = 8
        boss_chart.width = 18
        charts_sheet.add_chart(boss_chart, "A76")


def _add_chart_title(sheet, cell: str, title: str) -> None:
    sheet[cell] = title


def _add_table_if_possible(sheet, table_cls, table_style_cls) -> None:
    from openpyxl.utils import get_column_letter

    if sheet.max_row < 2 or sheet.max_column < 1 or sheet.title == "Charts":
        return
    if any(cell.value is None for cell in sheet[1]):
        return
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    if len(headers) != len(set(headers)) or any(not header for header in headers):
        return
    ref = "A1:%s%d" % (get_column_letter(sheet.max_column), sheet.max_row)
    table = table_cls(displayName=_safe_table_name(sheet.title), ref=ref)
    style = table_style_cls(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    try:
        sheet.add_table(table)
    except ValueError:
        return


def _style_sheet(sheet, font_cls, fill_cls, alignment_cls, border_cls, side_cls, column_letter_fn) -> None:
    title_font = font_cls(bold=True, size=15, color="FFD45E")
    header_fill = fill_cls("solid", fgColor="2D2140")
    header_font = font_cls(bold=True, color="FFD45E")
    thin_side = side_cls(style="thin", color="55436D")
    thin_border = border_cls(bottom=thin_side)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = alignment_cls(vertical="top", wrap_text=True)
    if sheet.max_row >= 1:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
    if sheet.title in {"Overview", "Learning Metrics"}:
        sheet["A1"].font = title_font
    for column_index in range(1, sheet.max_column + 1):
        column_letter = column_letter_fn(column_index)
        max_length = 12
        for cell in sheet[column_letter]:
            max_length = max(max_length, min(len(str(cell.value or "")), 64))
        sheet.column_dimensions[column_letter].width = min(max_length + 2, 70)
    sheet.freeze_panes = "A2"


def _research_questions() -> list[dict]:
    return [
        {
            "question": "Does performance improve over time?",
            "reading": "Use Timeline, Theme Metrics, success rate, and average time per attempt.",
        },
        {
            "question": "Where does the player struggle?",
            "reading": "Use Error Analysis, Interaction Metrics, deaths, and drop-off events.",
        },
        {
            "question": "Do hints and explanations help or create dependency?",
            "reading": "Use Support Impact and hint/explanation usage rates.",
        },
        {
            "question": "Does generated content work as well as patterns?",
            "reading": "Use AI Quality together with Validation Attempts and Generated Tasks.",
        },
        {
            "question": "Is the game engaging enough to support learning?",
            "reading": "Use session count, stage completions, terminal opens, and return-to-menu/exit events.",
        },
    ]


def _group_attempt_metrics(attempts: list[dict], events: list[dict], generated_tasks: list[dict], key: str) -> list[dict]:
    groups: dict[str, list[dict]] = {}
    for attempt in attempts:
        group_key = str(attempt.get(key, "") or "unknown")
        groups.setdefault(group_key, []).append(attempt)
    event_counts_by_group: dict[str, dict[str, int]] = {}
    for event in events:
        metadata = event.get("metadata", {})
        if not isinstance(metadata, dict):
            metadata = {}
        group_key = str(metadata.get(key, event.get(key, "")) or "unknown")
        event_type = str(event.get("event_type", "unknown"))
        event_counts_by_group.setdefault(group_key, {})
        event_counts_by_group[group_key][event_type] = event_counts_by_group[group_key].get(event_type, 0) + 1
    task_counts_by_group: dict[str, dict[str, int]] = {}
    for task in generated_tasks:
        group_key = str(task.get(key, "") or "unknown")
        source = str(task.get("generation_source", "unknown")).lower()
        task_counts_by_group.setdefault(group_key, {"ai": 0, "fallback": 0})
        if source == "ai":
            task_counts_by_group[group_key]["ai"] += 1
        if source == "fallback":
            task_counts_by_group[group_key]["fallback"] += 1
    all_keys = sorted(set(groups.keys()) | set(event_counts_by_group.keys()) | set(task_counts_by_group.keys()))
    rows: list[dict] = []
    for group_key in all_keys:
        group_attempts = groups.get(group_key, [])
        successes = sum(1 for attempt in group_attempts if attempt.get("success"))
        failures = len(group_attempts) - successes
        rows.append({
            key: group_key,
            "attempts": len(group_attempts),
            "successes": successes,
            "failures": failures,
            "success_rate_percent": _percent(successes, len(group_attempts)),
            "avg_time_seconds": _average([_safe_float(attempt.get("time_taken_seconds", 0)) for attempt in group_attempts]),
            "hint_used_attempts": sum(1 for attempt in group_attempts if bool(attempt.get("hint_used"))),
            "explanation_used_attempts": sum(1 for attempt in group_attempts if bool(attempt.get("explanation_used"))),
            "hint_events": event_counts_by_group.get(group_key, {}).get("hint_requested", 0),
            "explanation_events": event_counts_by_group.get(group_key, {}).get("explanation_opened", 0),
            "ai_tasks": task_counts_by_group.get(group_key, {}).get("ai", 0),
            "fallback_tasks": task_counts_by_group.get(group_key, {}).get("fallback", 0),
        })
    if not rows:
        rows.append({
            key: "no_data",
            "attempts": 0,
            "successes": 0,
            "failures": 0,
            "success_rate_percent": 0,
            "avg_time_seconds": 0,
            "hint_used_attempts": 0,
            "explanation_used_attempts": 0,
            "hint_events": 0,
            "explanation_events": 0,
            "ai_tasks": 0,
            "fallback_tasks": 0,
        })
    return rows


def _support_impact(attempts: list[dict]) -> list[dict]:
    groups = {
        "No support before submit": [attempt for attempt in attempts if not attempt.get("hint_used") and not attempt.get("explanation_used")],
        "Hint used": [attempt for attempt in attempts if attempt.get("hint_used")],
        "Explanation opened": [attempt for attempt in attempts if attempt.get("explanation_used")],
        "Hint and explanation": [attempt for attempt in attempts if attempt.get("hint_used") and attempt.get("explanation_used")],
    }
    rows: list[dict] = []
    for support_type, group_attempts in groups.items():
        successes = sum(1 for attempt in group_attempts if attempt.get("success"))
        rows.append({
            "support_type": support_type,
            "attempts": len(group_attempts),
            "successes": successes,
            "success_rate_percent": _percent(successes, len(group_attempts)),
            "avg_time_seconds": _average([_safe_float(attempt.get("time_taken_seconds", 0)) for attempt in group_attempts]),
        })
    return rows


def _error_analysis(attempts: list[dict]) -> list[dict]:
    counts: dict[str, int] = {}
    themes: dict[str, dict[str, int]] = {}
    interactions: dict[str, dict[str, int]] = {}
    for attempt in attempts:
        for error in _as_list(attempt.get("errors", [])):
            normalized = str(error).strip() or "unknown_error"
            counts[normalized] = counts.get(normalized, 0) + 1
            theme = str(attempt.get("level_theme", "unknown"))
            interaction = str(attempt.get("interaction_type", "unknown"))
            themes.setdefault(normalized, {})
            interactions.setdefault(normalized, {})
            themes[normalized][theme] = themes[normalized].get(theme, 0) + 1
            interactions[normalized][interaction] = interactions[normalized].get(interaction, 0) + 1
    rows = []
    for error, count in sorted(counts.items(), key=lambda item: item[1], reverse=True):
        rows.append({
            "error": error,
            "count": count,
            "most_common_theme": _top_key(themes.get(error, {})),
            "most_common_interaction": _top_key(interactions.get(error, {})),
        })
    if not rows:
        rows.append({"error": "no_errors_yet", "count": 0, "most_common_theme": "", "most_common_interaction": ""})
    return rows


def _ai_quality(generated_tasks: list[dict]) -> list[dict]:
    total = len(generated_tasks)
    buckets: dict[str, int] = {}
    detail_counts: dict[str, int] = {}
    for task in generated_tasks:
        source = str(task.get("generation_source", "") or task.get("storage_mode", "unknown"))
        if str(task.get("storage_mode", "")) == "prewarm_failed":
            source = "prewarm_failed"
        buckets[source] = buckets.get(source, 0) + 1
        detail = str(task.get("generation_detail", "") or task.get("detail", "") or "no_detail")
        detail_counts[detail] = detail_counts.get(detail, 0) + 1
    rows = [
        {
            "category": category,
            "count": count,
            "share_percent": _percent(count, total),
            "detail": "",
        }
        for category, count in sorted(buckets.items(), key=lambda item: item[1], reverse=True)
    ]
    for detail, count in sorted(detail_counts.items(), key=lambda item: item[1], reverse=True)[:10]:
        rows.append({
            "category": "detail",
            "count": count,
            "share_percent": _percent(count, total),
            "detail": detail,
        })
    if not rows:
        rows.append({"category": "no_generated_tasks_yet", "count": 0, "share_percent": 0, "detail": ""})
    return rows


def _terminal_tasks(attempts: list[dict]) -> list[dict]:
    grouped: dict[str, list[dict]] = {}
    for index, attempt in enumerate(attempts):
        session_id = str(attempt.get("terminal_session_id", "")).strip()
        if not session_id:
            session_id = "attempt_%04d" % index
        grouped.setdefault(session_id, []).append(attempt)
    rows: list[dict] = []
    for session_id, group_attempts in grouped.items():
        ordered = sorted(group_attempts, key=lambda item: (_safe_float(item.get("attempt_number", 0)), str(item.get("timestamp", ""))))
        first_attempt = ordered[0] if ordered else {}
        successful_attempts = [attempt for attempt in ordered if attempt.get("success")]
        success = bool(successful_attempts)
        success_attempt = successful_attempts[0] if successful_attempts else {}
        attempts_to_success = len(ordered)
        if success_attempt:
            attempt_number = int(_safe_float(success_attempt.get("attempt_number", 0)))
            if attempt_number > 0:
                attempts_to_success = attempt_number
        rows.append({
            "terminal_session_id": session_id,
            "level_theme": str(first_attempt.get("level_theme", "")),
            "interaction_type": str(first_attempt.get("interaction_type", "")),
            "attempts": attempts_to_success,
            "success": success,
            "first_try_success": success and attempts_to_success <= 1,
            "time_to_success_seconds": _safe_float(success_attempt.get("time_taken_seconds", 0)) if success_attempt else 0,
            "hint_used": any(bool(attempt.get("hint_used")) for attempt in ordered),
            "explanation_used": any(bool(attempt.get("explanation_used")) for attempt in ordered),
            "explanation_read_seconds": _average([_safe_float(attempt.get("explanation_read_seconds", 0)) for attempt in ordered]),
            "generation_source": str(first_attempt.get("generation_source", "")),
        })
    if not rows:
        rows.append({
            "terminal_session_id": "no_tasks_yet",
            "level_theme": "",
            "interaction_type": "",
            "attempts": 0,
            "success": False,
            "first_try_success": False,
            "time_to_success_seconds": 0,
            "hint_used": False,
            "explanation_used": False,
            "explanation_read_seconds": 0,
            "generation_source": "",
        })
    return rows


def _game_behavior(events: list[dict]) -> list[dict]:
    behavior_specs = [
        ("deaths", ["player_died", "death"], "Failure/frustration signal."),
        ("enemy_victories", ["enemy_defeated"], "Regular combat completion."),
        ("boss_victories", ["boss_defeated", "boss_completed"], "Boss-room completion."),
        ("player_damage_taken", ["player_damage_taken"], "Combat pressure on the player."),
        ("damage_dealt", ["player_damage_dealt", "boss_damage_dealt"], "How strongly code-enabled combat performs."),
        ("parry_attempts", ["parry_requested"], "How often the player tries to defend."),
        ("parry_successes", ["parry_succeeded", "projectile_reflected"], "Successful defensive skill usage."),
        ("parry_failures", ["parry_failed"], "Missed defensive timing."),
        ("terminal_closes_without_success", ["terminal_closed_without_success"], "Task abandonment or confusion."),
        ("combat_time_after_code", ["combat_unlocked"], "How long players can fight after coding before the next gate."),
    ]
    rows: list[dict] = []
    for behavior, event_names, interpretation in behavior_specs:
        matching = [event for event in events if str(event.get("event_type", "")) in event_names]
        values = [_metadata_number(event, "amount") for event in matching]
        if behavior == "combat_time_after_code":
            values = [_metadata_number(event, "elapsed_since_terminal_seconds") for event in matching]
        rows.append({
            "behavior": behavior,
            "count": len(matching),
            "total_value": round(sum(values), 2),
            "average_value": _average(values),
            "interpretation": interpretation,
        })
    return rows


def _difficulty_progression(events: list[dict]) -> list[dict]:
    rows: list[dict] = []
    for event in events:
        if str(event.get("event_type", "")) != "adaptive_difficulty_changed":
            continue
        metadata = event.get("metadata", {})
        if not isinstance(metadata, dict):
            metadata = {}
        rows.append({
            "timestamp": event.get("timestamp", ""),
            "interaction_type": metadata.get("interaction_type", ""),
            "level_theme": metadata.get("level_theme", event.get("level_id", "")),
            "from_difficulty": metadata.get("from_difficulty", ""),
            "to_difficulty": metadata.get("to_difficulty", ""),
            "reason": metadata.get("reason", ""),
        })
    if not rows:
        rows.append({"timestamp": "", "interaction_type": "", "level_theme": "", "from_difficulty": "", "to_difficulty": "", "reason": "no_changes_yet"})
    return rows


def _progression_funnel(events: list[dict]) -> list[dict]:
    stages = [
        ("tutorial", "tutorial_started", "tutorial_completed"),
        ("level_01", "level_started", "level_completed"),
        ("boss_01", "boss_started", "boss_completed"),
        ("level_02", "level_started", "level_completed"),
        ("boss_02", "boss_started", "boss_completed"),
        ("level_03", "level_started", "level_completed"),
        ("boss_03", "boss_started", "boss_completed"),
        ("level_04", "level_started", "level_completed"),
        ("boss_04", "boss_started", "boss_completed"),
        ("level_05", "level_started", "level_completed"),
        ("boss_05", "boss_started", "boss_completed"),
    ]
    rows: list[dict] = []
    for stage, start_event, complete_event in stages:
        started = _stage_event_count(events, start_event, stage)
        completed = _stage_event_count(events, complete_event, stage)
        rows.append({
            "stage": stage,
            "started": started,
            "completed": completed,
            "completion_rate_percent": _percent(completed, started),
            "dropoff_count": max(started - completed, 0),
        })
    return rows


def _boss_performance(attempts: list[dict], events: list[dict]) -> list[dict]:
    level_ids = sorted({str(event.get("level_id", "")) for event in events if str(event.get("stage_type", "")) == "boss" or str(event.get("event_type", "")).startswith("boss")})
    if not level_ids:
        level_ids = ["no_boss_data_yet"]
    rows: list[dict] = []
    for level_id in level_ids:
        boss_events = [event for event in events if str(event.get("level_id", "")) == level_id]
        boss_attempts = [attempt for attempt in attempts if str(attempt.get("interaction_type", "")) == "boss" and _attempt_level_id(attempt, events) == level_id]
        rows.append({
            "level_id": level_id,
            "attempts": len(boss_attempts),
            "victories": sum(1 for event in boss_events if event.get("event_type") in {"boss_defeated", "boss_completed"}),
            "deaths": sum(1 for event in boss_events if event.get("event_type") in {"player_died", "death"}),
            "terminal_rewrites": sum(1 for event in boss_events if event.get("event_type") == "terminal_opened"),
            "player_damage_taken": round(sum(_metadata_number(event, "amount") for event in boss_events if event.get("event_type") == "player_damage_taken"), 2),
            "damage_dealt_to_boss": round(sum(_metadata_number(event, "amount") for event in boss_events if event.get("event_type") == "boss_damage_dealt"), 2),
            "parry_successes": sum(1 for event in boss_events if event.get("event_type") in {"parry_succeeded", "projectile_reflected"}),
        })
    return rows


def _retention(events: list[dict]) -> list[dict]:
    by_date: dict[str, dict[str, float | set[str]]] = {}
    session_start_times: dict[str, int] = {}
    session_end_times: dict[str, int] = {}
    for event in events:
        date_key = str(event.get("timestamp", ""))[:10] or "unknown"
        bucket = by_date.setdefault(date_key, {"sessions": set(), "events": 0, "duration": 0.0})
        session_id = str(event.get("session_id", ""))
        if session_id:
            bucket["sessions"].add(session_id)  # type: ignore[union-attr]
        bucket["events"] = float(bucket["events"]) + 1
        event_type = str(event.get("event_type", ""))
        timestamp_unix = int(_safe_float(event.get("timestamp_unix", 0)))
        if event_type == "session_start" and session_id:
            session_start_times[session_id] = timestamp_unix
        if event_type == "session_end" and session_id:
            session_end_times[session_id] = timestamp_unix
    for session_id, start_time in session_start_times.items():
        end_time = session_end_times.get(session_id, start_time)
        date_key = datetime.fromtimestamp(start_time, timezone.utc).strftime("%Y-%m-%d")
        bucket = by_date.setdefault(date_key, {"sessions": set(), "events": 0, "duration": 0.0})
        bucket["duration"] = float(bucket["duration"]) + max(end_time - start_time, 0)
    seen_sessions: set[str] = set()
    rows: list[dict] = []
    for date_key in sorted(by_date.keys()):
        sessions = by_date[date_key]["sessions"]
        session_set = sessions if isinstance(sessions, set) else set()
        returning = len(session_set.intersection(seen_sessions))
        seen_sessions.update(session_set)
        rows.append({
            "date_utc": date_key,
            "sessions": len(session_set),
            "events": int(float(by_date[date_key]["events"])),
            "returning_session_count": returning,
            "session_duration_seconds": round(float(by_date[date_key]["duration"]), 1),
        })
    if not rows:
        rows.append({"date_utc": "no_sessions_yet", "sessions": 0, "events": 0, "returning_session_count": 0, "session_duration_seconds": 0})
    return rows


def _timeline(events: list[dict], attempts: list[dict]) -> list[dict]:
    buckets: dict[str, dict[str, int]] = {}
    for event in events:
        bucket = _hour_bucket(str(event.get("timestamp", "")))
        buckets.setdefault(bucket, _empty_timeline_bucket())
        buckets[bucket]["events"] += 1
        event_type = str(event.get("event_type", ""))
        if event_type == "hint_requested":
            buckets[bucket]["hint_requests"] += 1
        if event_type == "explanation_opened":
            buckets[bucket]["explanation_opens"] += 1
    for attempt in attempts:
        bucket = _hour_bucket(str(attempt.get("timestamp", "")))
        buckets.setdefault(bucket, _empty_timeline_bucket())
        buckets[bucket]["validation_attempts"] += 1
        if attempt.get("success"):
            buckets[bucket]["successful_validations"] += 1
        else:
            buckets[bucket]["failed_validations"] += 1
    rows = [{"hour_utc": bucket, **buckets[bucket]} for bucket in sorted(buckets.keys())]
    if not rows:
        rows.append({"hour_utc": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:00"), **_empty_timeline_bucket()})
    return rows


def _event_counts(events: list[dict]) -> list[dict]:
    counts: dict[str, int] = {}
    for event in events:
        event_type = str(event.get("event_type", "unknown"))
        counts[event_type] = counts.get(event_type, 0) + 1
    rows = [{"event_type": event_type, "count": count} for event_type, count in sorted(counts.items(), key=lambda item: item[1], reverse=True)]
    if not rows:
        rows.append({"event_type": "no_events_yet", "count": 0})
    return rows


def _build_conclusions(learning_metrics: list[dict], attempts: list[dict], events: list[dict], generated_tasks: list[dict]) -> list[dict]:
    metric_map = {str(item["metric"]): item["value"] for item in learning_metrics}
    success_rate = _safe_float(metric_map.get("Validation success rate, %", 0))
    hint_rate = _safe_float(metric_map.get("Attempts after hint, %", 0))
    explanation_rate = _safe_float(metric_map.get("Attempts after explanation, %", 0))
    fallback_count = _safe_float(metric_map.get("Fallback pattern tasks", 0))
    ai_count = _safe_float(metric_map.get("AI-generated tasks", 0))
    ai_total = ai_count + fallback_count
    first_half, second_half = _split_attempts(attempts)
    first_success = _percent(sum(1 for attempt in first_half if attempt.get("success")), len(first_half))
    second_success = _percent(sum(1 for attempt in second_half if attempt.get("success")), len(second_half))
    conclusions = [
        {
            "conclusion": _learning_signal_text(success_rate, first_success, second_success),
            "signal": "Overall success %s%%; first half %s%%; second half %s%%." % (success_rate, first_success, second_success),
            "recommended_action": "Keep comparing early and late attempts for the same theme before claiming learning impact.",
        },
        {
            "conclusion": "Support demand is %s." % ("high" if hint_rate >= 50 or explanation_rate >= 50 else "moderate/low"),
            "signal": "Hints used in %s%% of attempts; explanations in %s%%." % (hint_rate, explanation_rate),
            "recommended_action": "If support is high and success rises, explanations help. If support is high and success stays low, tasks are too hard or unclear.",
        },
        {
            "conclusion": "AI content availability is %s." % ("stable" if ai_total and _percent(ai_count, int(ai_total)) >= 50 else "limited"),
            "signal": "AI tasks: %s; fallback tasks: %s; total generated records: %s." % (int(ai_count), int(fallback_count), len(generated_tasks)),
            "recommended_action": "Use AI Quality to separate learning problems from provider/model access problems.",
        },
        {
            "conclusion": "Engagement should be read together with learning results.",
            "signal": "Events: %s; terminal opens: %s; stage completions: %s." % (
                len(events),
                _count_events(events, "terminal_opened"),
                sum(1 for event in events if str(event.get("event_type", "")).endswith("_completed")),
            ),
            "recommended_action": "Positive impact means both retention/progression and concept success improve.",
        },
    ]
    return conclusions


def _learning_signal_text(success_rate: float, first_success: float, second_success: float) -> str:
    if second_success > first_success + 10:
        return "Learning trend looks positive."
    if success_rate < 35:
        return "Learning trend is weak or tasks are too difficult."
    if second_success < first_success - 10:
        return "Learning trend may be negative or fatigue/friction is high."
    return "Learning trend is mixed; collect more attempts."


def _empty_timeline_bucket() -> dict[str, int]:
    return {
        "events": 0,
        "validation_attempts": 0,
        "successful_validations": 0,
        "failed_validations": 0,
        "hint_requests": 0,
        "explanation_opens": 0,
    }


def _metadata_number(event: dict, key: str) -> float:
    metadata = event.get("metadata", {})
    if not isinstance(metadata, dict):
        return 0.0
    return _safe_float(metadata.get(key, 0))


def _stage_event_count(events: list[dict], event_type: str, stage: str) -> int:
    count = 0
    for event in events:
        if event.get("event_type") != event_type:
            continue
        if stage == "tutorial":
            count += 1
            continue
        level_number = "".join(char for char in stage if char.isdigit())
        level_id = str(event.get("level_id", ""))
        if level_number and level_id.endswith(level_number):
            count += 1
    return count


def _attempt_level_id(attempt: dict, _events: list[dict]) -> str:
    return str(attempt.get("level_id", ""))


def _count_events(events: list[dict], event_type: str) -> int:
    return sum(1 for event in events if event.get("event_type") == event_type)


def _percent(numerator: int | float, denominator: int | float) -> float:
    return round(float(numerator) / float(denominator) * 100.0, 1) if denominator else 0.0


def _average(values: list[float]) -> float:
    clean_values = [value for value in values if value >= 0]
    return round(sum(clean_values) / len(clean_values), 1) if clean_values else 0.0


def _safe_float(value) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _as_list(value) -> list:
    if isinstance(value, list):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            return [value] if value.strip() else []
        return parsed if isinstance(parsed, list) else [parsed]
    return [value] if value else []


def _top_key(counts: dict[str, int]) -> str:
    if not counts:
        return ""
    return max(counts.items(), key=lambda item: item[1])[0]


def _split_attempts(attempts: list[dict]) -> tuple[list[dict], list[dict]]:
    if len(attempts) < 2:
        return attempts, attempts
    sorted_attempts = sorted(attempts, key=lambda item: str(item.get("timestamp", "")))
    midpoint = max(1, len(sorted_attempts) // 2)
    return sorted_attempts[:midpoint], sorted_attempts[midpoint:]


def _safe_table_name(sheet_title: str) -> str:
    cleaned = "".join(char if char.isalnum() else "_" for char in sheet_title)
    cleaned = cleaned.strip("_") or "Sheet"
    return "Table_%s" % cleaned[:24]


def _cell_value(value) -> str | int | float | bool:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return "" if value is None else value
    return json.dumps(value, ensure_ascii=False)


def _hour_bucket(timestamp: str) -> str:
    try:
        parsed = datetime.fromisoformat(timestamp.replace("Z", "+00:00"))
    except ValueError:
        parsed = datetime.now(timezone.utc)
    return parsed.astimezone(timezone.utc).strftime("%Y-%m-%d %H:00")
